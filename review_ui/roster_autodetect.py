"""
Roster column autodetection with confidence scoring
Intelligently maps Excel/CSV columns to patient fields
"""

import re
import openpyxl
import pandas as pd
from datetime import datetime
from typing import Dict, List, Tuple, Optional

COLUMN_PATTERNS = {
    'last': [
        r'(?i)\b(last|surname|l\s*name|family|last_name|lastname)\b'
    ],
    'first': [
        r'(?i)\b(first|given|f\s*name|first_name|firstname|given_name)\b'
    ],
    'dob': [
        r'(?i)\b(dob|date.?of.?birth|birthdate|birth.?date|birth|bdate|d\.?o\.?b\.?)\b'
    ]
}

def score_column_name(column_name: str, field: str) -> float:
    """
    Score how well a column name matches a field
    Returns confidence score 0.0-1.0
    """
    if not column_name:
        return 0.0
    
    column_clean = column_name.strip().lower()
    patterns = COLUMN_PATTERNS.get(field, [])
    
    for pattern in patterns:
        match = re.search(pattern, column_clean)
        if match:
            if match.group(0) == column_clean:
                return 1.0
            else:
                return 0.9
    
    return 0.0


def score_column_values(values: List, field: str) -> float:
    """
    Score how well column values match expected data type
    Returns confidence score 0.0-1.0
    """
    if not values or len(values) == 0:
        return 0.0
    
    values = [v for v in values if v is not None and str(v).strip()]
    if len(values) == 0:
        return 0.0
    
    if field in ['last', 'first']:
        name_pattern = r'^[A-Za-z\'\-\s]+$'
        matches = sum(1 for v in values if re.match(name_pattern, str(v)))
        return matches / len(values)
    
    elif field == 'dob':
        date_matches = 0
        for v in values:
            if isinstance(v, datetime):
                date_matches += 1
            else:
                v_str = str(v).strip()
                if re.match(r'\d{4}-\d{2}-\d{2}', v_str):
                    date_matches += 1
                elif re.match(r'\d{1,2}[/-]\d{1,2}[/-]\d{2,4}', v_str):
                    date_matches += 1
        
        return date_matches / len(values)
    
    return 0.0


def autodetect_columns(file_path: str, max_rows: int = 50) -> Tuple[Dict[str, str], Dict[str, float]]:
    """
    Auto-detect roster column mapping
    Returns (detected_columns, confidence_scores)
    """
    ext = file_path.lower().split('.')[-1]
    
    if ext in ['xlsx', 'xls']:
        return autodetect_excel(file_path, max_rows)
    elif ext == 'csv':
        return autodetect_csv(file_path, max_rows)
    elif ext == 'json':
        return autodetect_json(file_path)
    else:
        return {}, {}


def autodetect_excel(file_path: str, max_rows: int = 50) -> Tuple[Dict[str, str], Dict[str, float]]:
    """Auto-detect columns in Excel file"""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    
    headers = []
    data_rows = []
    
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_idx == 1:
            headers = [str(cell).strip() if cell else f'Column_{i}' for i, cell in enumerate(row)]
            continue
        
        if row_idx > max_rows:
            break
        
        if any(row):
            data_rows.append(row)
    
    wb.close()
    
    return detect_from_headers_and_data(headers, data_rows)


def autodetect_csv(file_path: str, max_rows: int = 50) -> Tuple[Dict[str, str], Dict[str, float]]:
    """Auto-detect columns in CSV file"""
    df = pd.read_csv(file_path, nrows=max_rows)
    headers = df.columns.tolist()
    data_rows = df.values.tolist()
    
    return detect_from_headers_and_data(headers, data_rows)


def autodetect_json(file_path: str) -> Tuple[Dict[str, str], Dict[str, float]]:
    """Auto-detect columns in JSON file"""
    import json
    with open(file_path, 'r') as f:
        data = json.load(f)
    
    if isinstance(data, list) and len(data) > 0:
        headers = list(data[0].keys())
        data_rows = [[item.get(k) for k in headers] for item in data[:50]]
        return detect_from_headers_and_data(headers, data_rows)
    
    return {}, {}


def detect_from_headers_and_data(headers: List[str], data_rows: List[List]) -> Tuple[Dict[str, str], Dict[str, float]]:
    """
    Detect columns from headers and sample data
    Returns (detected_columns, confidence_scores)
    """
    detected = {}
    confidence = {}
    
    for field in ['last', 'first', 'dob']:
        best_column = None
        best_score = 0.0
        
        for col_idx, header in enumerate(headers):
            name_score = score_column_name(header, field)
            
            column_values = [row[col_idx] if col_idx < len(row) else None for row in data_rows]
            value_score = score_column_values(column_values, field)
            
            combined_score = (name_score * 0.7) + (value_score * 0.3)
            
            if combined_score > best_score:
                best_score = combined_score
                best_column = header
        
        if best_score > 0.5:
            detected[field] = best_column
            confidence[field] = round(best_score, 2)
    
    return detected, confidence


def get_sample_data(file_path: str, detected_columns: Dict[str, str], limit: int = 5) -> List[Dict]:
    """
    Get sample patient data using detected columns
    """
    ext = file_path.lower().split('.')[-1]
    
    if ext in ['xlsx', 'xls']:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        
        rows_list = list(ws.iter_rows(values_only=True))
        if len(rows_list) < 2:
            return []
        
        headers = [str(cell).strip() if cell else f'Column_{i}' for i, cell in enumerate(rows_list[0])]
        
        samples = []
        for row in rows_list[1:limit+1]:
            sample = {}
            for field, column_name in detected_columns.items():
                if column_name in headers:
                    col_idx = headers.index(column_name)
                    if col_idx < len(row):
                        value = row[col_idx]
                        if value is not None:
                            if field == 'dob' and isinstance(value, datetime):
                                sample[field] = value.strftime('%Y-%m-%d')
                            else:
                                sample[field] = str(value)
            
            if sample:
                samples.append(sample)
        
        wb.close()
        return samples
    
    elif ext == 'csv':
        df = pd.read_csv(file_path, nrows=limit)
        samples = []
        for _, row in df.iterrows():
            sample = {}
            for field, column_name in detected_columns.items():
                if column_name in df.columns:
                    value = row[column_name]
                    if pd.notna(value):
                        sample[field] = str(value)
            if sample:
                samples.append(sample)
        return samples
    
    return []
