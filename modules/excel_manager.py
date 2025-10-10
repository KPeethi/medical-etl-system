"""
Excel Template Generator and Log Exporter for Medical ETL System
Creates Excel templates and exports logs to Excel format
"""

import pandas as pd
from pathlib import Path
from datetime import datetime
import json
from typing import Dict, List, Any, Optional
from config.config import Config

class ExcelManager:
    """Manages Excel templates and log exports"""
    
    def __init__(self):
        self.config = Config()
        
    def create_mapping_template(self, output_path: Path, include_examples: bool = True) -> bool:
        """
        Create Excel mapping template with proper formatting
        
        Args:
            output_path: Path where to save the template
            include_examples: Whether to include example data
            
        Returns:
            bool: True if successful
        """
        try:
            # Define template structure
            if include_examples:
                data = {
                    'Patient_ID': [12345, 67890, 11111, 22222, 33333, 44444, 55555],
                    'Last_Name': ['Smith', 'Johnson', 'Brown', 'Davis', 'Wilson', 'Garcia', 'Martinez'],
                    'First_Name': ['John', 'Mary', 'David', 'Sarah', 'Michael', 'Lisa', 'Robert'],
                    'DOB': ['01-15-1980', '03-22-1975', '05-10-1990', '12-03-1985', '08-20-1970', '02-14-1995', '11-30-1988'],
                    'File_Name': [
                        'john_smith_chart.pdf',
                        'mary_johnson_xray.jpg', 
                        'david_brown_lab.pdf',
                        'sarah_davis_report.png',
                        'michael_wilson_scan.tiff',
                        'lisa_garcia_document.pdf',
                        'robert_martinez_image.jpg'
                    ]
                }
            else:
                # Empty template with headers only
                data = {
                    'Patient_ID': [],
                    'Last_Name': [],
                    'First_Name': [],
                    'DOB': [],
                    'File_Name': []
                }
            
            df = pd.DataFrame(data)
            
            # Create Excel file with formatting
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                # Main mapping sheet
                df.to_excel(writer, sheet_name='Patient_Mapping', index=False)
                
                # Get workbook and worksheet
                workbook = writer.book
                worksheet = writer.sheets['Patient_Mapping']
                
                # Format headers
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                
                header_font = Font(bold=True, color='FFFFFF')
                header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
                header_alignment = Alignment(horizontal='center', vertical='center')
                
                # Apply header formatting
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Create instructions sheet
                instructions_data = {
                    'Column Name': [
                        'Patient_ID',
                        'Last_Name', 
                        'First_Name',
                        'DOB',
                        'File_Name'
                    ],
                    'Description': [
                        'Unique patient identifier (numbers/text)',
                        'Patient last name/surname', 
                        'Patient first name/given name',
                        'Date of birth (MM-DD-YYYY format)',
                        'Original filename (with extension)'
                    ],
                    'Required': [
                        'Yes (or File_Name)',
                        'Yes',
                        'Yes', 
                        'Optional',
                        'Yes (or Patient_ID)'
                    ],
                    'Examples': [
                        '12345, H12345, MRN-67890',
                        'Smith, Johnson, García',
                        'John, Mary, José',
                        '01-15-1980, 12-03-1985',
                        'patient_chart.pdf, scan.jpg'
                    ]
                }
                
                instructions_df = pd.DataFrame(instructions_data)
                instructions_df.to_excel(writer, sheet_name='Instructions', index=False)
                
                # Format instructions sheet
                instructions_ws = writer.sheets['Instructions']
                for cell in instructions_ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                
                # Auto-adjust instruction columns
                for column in instructions_ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 60)
                    instructions_ws.column_dimensions[column_letter].width = adjusted_width
                
                # Add field synonyms sheet
                synonyms_data = {
                    'Field Type': ['Patient ID'] * 12 + ['Last Name'] * 9 + ['First Name'] * 9 + ['File Name'] * 9 + ['Date of Birth'] * 6,
                    'Accepted Column Names': [
                        'id', 'patient_id', 'patientid', 'pno', 'p_no', 'patient_no',
                        'patient_number', 'id_number', 'idnumber', 'entity_id', 'medical_id', 'mrn',
                        'lastname', 'last_name', 'lname', 'surname', 'family_name', 'familyname', 'last', 'sur_name', 'l_name',
                        'firstname', 'first_name', 'fname', 'givenname', 'given_name', 'first', 'forename', 'christian_name', 'f_name',
                        'filename', 'file_name', 'file', 'document', 'doc', 'document_name', 'doc_name', 'name', 'filepath',
                        'dob', 'date_of_birth', 'dateofbirth', 'birth_date', 'birthdate', 'birth'
                    ]
                }
                
                synonyms_df = pd.DataFrame(synonyms_data)
                synonyms_df.to_excel(writer, sheet_name='Field_Synonyms', index=False)
                
                # Format synonyms sheet
                synonyms_ws = writer.sheets['Field_Synonyms']
                for cell in synonyms_ws[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                
                for column in synonyms_ws.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 40)
                    synonyms_ws.column_dimensions[column_letter].width = adjusted_width
            
            return True
            
        except Exception as e:
            print(f"Error creating Excel template: {str(e)}")
            return False
    
    def export_session_to_excel(self, session_summary_path: Path, output_path: Path) -> bool:
        """
        Export session logs to Excel format
        
        Args:
            session_summary_path: Path to JSON session summary
            output_path: Path where to save Excel log
            
        Returns:
            bool: True if successful
        """
        try:
            # Load session data
            with open(session_summary_path, 'r', encoding='utf-8') as f:
                session_data = json.load(f)
            
            # Create Excel file with multiple sheets
            with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
                
                # 1. SESSION SUMMARY SHEET
                summary = session_data.get('summary', {})
                detailed_stats = session_data.get('detailed_stats', {})
                
                summary_data = {
                    'Metric': [
                        'Session ID', 'Run Type', 'Duration (minutes)', 'Total Files Processed',
                        'Successful Operations', 'Failed Operations', 'Duplicate Files Found',
                        'Unmapped Files', 'Unique Patients Processed', 'Total Errors', 'Total Warnings'
                    ],
                    'Value': [
                        summary.get('session_id', ''),
                        summary.get('run_type', ''),
                        round(summary.get('duration_seconds', 0) / 60, 2),
                        summary.get('total_files_processed', 0),
                        summary.get('successful_operations', 0),
                        summary.get('failed_operations', 0),
                        summary.get('duplicate_files_found', 0),
                        summary.get('unmapped_files', 0),
                        summary.get('unique_patients_processed', 0),
                        summary.get('total_errors', 0),
                        summary.get('total_warnings', 0)
                    ]
                }
                
                summary_df = pd.DataFrame(summary_data)
                summary_df.to_excel(writer, sheet_name='Session_Summary', index=False)
                
                # 2. FILE OPERATIONS SHEET (with source and destination paths)
                file_operations = session_data.get('file_operations', [])
                
                if file_operations:
                    file_data = []
                    for op in file_operations:
                        patient_data = op.get('patient_data', {})
                        file_data.append({
                            'Operation_ID': op.get('operation_id', ''),
                            'Source_File': op.get('source_file', ''),
                            'Destination_File': op.get('destination_file', ''),
                            'File_Size_MB': round(op.get('file_size', 0) / (1024*1024), 2),
                            'File_Type': op.get('file_type', ''),
                            'Status': op.get('status', ''),
                            'Patient_ID': patient_data.get('id', ''),
                            'Patient_First_Name': patient_data.get('firstname', ''),
                            'Patient_Last_Name': patient_data.get('lastname', ''),
                            'Patient_DOB': patient_data.get('dob', ''),
                            'Data_Source': op.get('data_source', ''),
                            'Processing_Duration_Sec': round((datetime.fromisoformat(op.get('end_time', '2025-01-01T00:00:00')) - 
                                                           datetime.fromisoformat(op.get('start_time', '2025-01-01T00:00:00'))).total_seconds(), 2) if op.get('end_time') and op.get('start_time') else 0,
                            'Errors': ', '.join(op.get('errors', [])),
                            'Warnings': ', '.join(op.get('warnings', []))
                        })
                    
                    files_df = pd.DataFrame(file_data)
                    files_df.to_excel(writer, sheet_name='File_Operations', index=False)
                
                # 3. DATA SOURCE BREAKDOWN SHEET
                data_sources = detailed_stats.get('data_sources', {})
                if data_sources:
                    ds_data = {
                        'Data_Source': list(data_sources.keys()),
                        'File_Count': list(data_sources.values()),
                        'Percentage': [round((count / sum(data_sources.values())) * 100, 1) 
                                     for count in data_sources.values()]
                    }
                    
                    ds_df = pd.DataFrame(ds_data)
                    ds_df.to_excel(writer, sheet_name='Data_Sources', index=False)
                
                # 4. DEEP SCAN STATISTICS SHEET
                deep_scan_stats = detailed_stats.get('deep_scan_statistics', {})
                if deep_scan_stats:
                    deep_scan_data = {
                        'Metric': [
                            'Files Requiring Deep Scan',
                            'Deep Scan Successful',
                            'Deep Scan Failed', 
                            'Average Pages Processed',
                            'Total Pages Processed'
                        ],
                        'Value': [
                            deep_scan_stats.get('files_requiring_deep_scan', 0),
                            deep_scan_stats.get('deep_scan_successful', 0),
                            deep_scan_stats.get('deep_scan_failed', 0),
                            deep_scan_stats.get('average_pages_processed', 0),
                            deep_scan_stats.get('total_pages_processed', 0)
                        ]
                    }
                    
                    deep_scan_df = pd.DataFrame(deep_scan_data)
                    deep_scan_df.to_excel(writer, sheet_name='Deep_Scan_Stats', index=False)
                
                # 5. ERRORS AND WARNINGS SHEET
                errors = detailed_stats.get('errors', [])
                warnings = detailed_stats.get('warnings', [])
                
                issues_data = []
                for error in errors:
                    issues_data.append({
                        'Type': 'ERROR',
                        'Operation': error.get('operation', ''),
                        'Message': error.get('error', ''),
                        'Timestamp': error.get('timestamp', ''),
                        'Context': str(error.get('context', {}))
                    })
                
                for warning in warnings:
                    issues_data.append({
                        'Type': 'WARNING',
                        'Operation': warning.get('operation', ''),
                        'Message': warning.get('warning', ''),
                        'Timestamp': warning.get('timestamp', ''),
                        'Context': str(warning.get('context', {}))
                    })
                
                if issues_data:
                    issues_df = pd.DataFrame(issues_data)
                    issues_df.to_excel(writer, sheet_name='Errors_Warnings', index=False)
                
                # Apply formatting to all sheets
                self._format_excel_sheets(writer)
            
            return True
            
        except Exception as e:
            print(f"Error exporting session to Excel: {str(e)}")
            return False
    
    def _format_excel_sheets(self, writer):
        """Apply consistent formatting to Excel sheets"""
        try:
            from openpyxl.styles import Font, PatternFill, Alignment
            
            header_font = Font(bold=True, color='FFFFFF')
            header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
            header_alignment = Alignment(horizontal='center', vertical='center')
            
            for sheet_name in writer.sheets:
                worksheet = writer.sheets[sheet_name]
                
                # Format headers
                for cell in worksheet[1]:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = min(max_length + 2, 60)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                    
        except Exception as e:
            print(f"Error formatting Excel sheets: {str(e)}")


def create_sample_files():
    """Create sample Excel mapping and log files"""
    excel_manager = ExcelManager()
    
    # Create mapping template with examples
    config = Config()
    template_path = config.DATA_DIR / "Patient_Mapping_Template.xlsx"
    
    success = excel_manager.create_mapping_template(template_path, include_examples=True)
    if success:
        print(f"Excel mapping template created: {template_path}")
    
    # Create empty template
    empty_template_path = config.DATA_DIR / "Patient_Mapping_Empty.xlsx"
    success = excel_manager.create_mapping_template(empty_template_path, include_examples=False)
    if success:
        print(f"Empty Excel mapping template created: {empty_template_path}")


if __name__ == "__main__":
    create_sample_files()