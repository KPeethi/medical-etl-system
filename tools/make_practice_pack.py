#!/usr/bin/env python3
"""
Practice Pack Generator
Creates a complete practice pack template for easy onboarding
"""

import os
import sys
import argparse
from pathlib import Path
import openpyxl
import yaml


PRACTICE_PACK_README = """# {practice_name} Practice Pack

## Quick Start

1. **Update patient roster**: Edit `patients.xlsx` with your actual patient list
2. **Review configuration**: Check `config.yml` for any practice-specific settings
3. **Test with samples**: Run a dry-run on the `samples/` folder first
4. **Process real data**: Use the Postman request below

## Postman Request

```json
POST http://localhost:5000/api/process
{{
  "source": "C:/path/to/your/files",
  "dest": "C:/output",
  "dry_run": true,
  "mapping": {{
    "identity": {{
      "roster": {{
        "path": "{pack_path}/patients.xlsx"
      }}
    }}
  }}
}}
```

## Files in This Pack

- `patients.xlsx` - Patient roster (REQUIRED - edit this!)
- `config.yml` - Practice-specific configuration
- `samples/` - Sample files for testing
- `README.txt` - This file

## Support

If files go to "unmapped", check:
1. Roster file has correct columns: last_name, first_name, dob
2. Patient names in files match roster
3. Use GET /api/identity/preview to test roster detection
"""


DEFAULT_CONFIG = """# {practice_name} Configuration

practice_id: "{practice_id}"

naming:
  patient_folder: "{{Last}}, {{First}} {{DOB_MM}}-{{DOB_DD}}-{{DOB_YYYY}}"
  strip_middle_names: true

identity:
  roster:
    type: excel
    path: "{pack_path}/patients.xlsx"
    autodetect_headers: true
    hints:
      last: ["last", "last_name", "surname", "family_name"]
      first: ["first", "first_name", "given", "given_name"]
      dob: ["dob", "date_of_birth", "birth", "birthdate"]

module_detection:
  folder_keywords:
    labs: ["lab", "labs", "laboratory", "pathology"]
    imaging: ["xray", "x-ray", "ct", "mri", "imaging", "radiology"]
    notes: ["note", "notes", "progress", "encounter", "clinical"]
    reports: ["report", "reports", "summary"]
  filename_patterns:
    labs: [".*\\\\bLAB\\\\b.*", ".*\\\\bPATH\\\\b.*"]
    imaging: [".*\\\\bXRAY\\\\b.*", ".*\\\\bCT\\\\b.*", ".*\\\\bMRI\\\\b.*"]

dedupe_policy: "same_size_same_module"

include_exts: [".pdf", ".tif", ".tiff", ".jpg", ".jpeg", ".png", ".doc", ".docx", ".rtf", ".txt", ".xlsx", ".csv"]

exclude_exts: [".exe", ".dll", ".zip", ".7z", ".rar"]

max_depth: 12

safety:
  source_readonly_check: true
  dry_run_default: true

logging:
  csv_dir: "./logs"
  redaction: true
  enable_db_logging: true
"""


def create_practice_pack(name: str, output_dir: str):
    """Create a complete practice pack"""
    pack_dir = Path(output_dir) / name
    pack_dir.mkdir(parents=True, exist_ok=True)
    
    practice_id = name.replace(' ', '_').replace('-', '_')
    
    print(f"Creating practice pack: {name}")
    print(f"Output directory: {pack_dir}")
    
    create_patient_roster(pack_dir / "patients.xlsx")
    print("✓ Created patients.xlsx template")
    
    config_content = DEFAULT_CONFIG.format(
        practice_name=name,
        practice_id=practice_id,
        pack_path=str(pack_dir).replace('\\', '/')
    )
    with open(pack_dir / "config.yml", 'w') as f:
        f.write(config_content)
    print("✓ Created config.yml")
    
    readme_content = PRACTICE_PACK_README.format(
        practice_name=name,
        pack_path=str(pack_dir).replace('\\', '/')
    )
    with open(pack_dir / "README.txt", 'w') as f:
        f.write(readme_content)
    print("✓ Created README.txt")
    
    samples_dir = pack_dir / "samples"
    samples_dir.mkdir(exist_ok=True)
    
    create_sample_file(samples_dir / "Smith_John_1980-05-15_lab.txt", 
                       "Sample lab report for John Smith")
    create_sample_file(samples_dir / "Johnson_Mary_1975-03-20_xray.txt", 
                       "Sample X-ray report for Mary Johnson")
    print(f"✓ Created samples/ directory with 2 test files")
    
    print(f"\n✅ Practice pack created successfully!")
    print(f"\nNext steps:")
    print(f"1. Edit {pack_dir}/patients.xlsx with actual patient data")
    print(f"2. Review {pack_dir}/config.yml")
    print(f"3. Test with: POST /api/process pointing to this pack")


def create_patient_roster(file_path: Path):
    """Create sample patient roster Excel file"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Patients"
    
    ws['A1'] = 'last_name'
    ws['B1'] = 'first_name'
    ws['C1'] = 'dob'
    
    ws['A2'] = 'Smith'
    ws['B2'] = 'John'
    ws['C2'] = '1980-05-15'
    
    ws['A3'] = 'Johnson'
    ws['B3'] = 'Mary'
    ws['C3'] = '1975-03-20'
    
    ws['A4'] = 'Brown'
    ws['B4'] = 'Patricia'
    ws['C4'] = '1985-03-10'
    
    for col in ['A', 'B', 'C']:
        ws.column_dimensions[col].width = 15
    
    wb.save(file_path)


def create_sample_file(file_path: Path, content: str):
    """Create a sample test file"""
    with open(file_path, 'w') as f:
        f.write(content)


def main():
    parser = argparse.ArgumentParser(description='Generate Medical ETL Practice Pack')
    parser.add_argument('--name', required=True, help='Practice name (e.g., "Alexander OBGYN")')
    parser.add_argument('--out', default='.', help='Output directory (default: current directory)')
    
    args = parser.parse_args()
    
    create_practice_pack(args.name, args.out)


if __name__ == '__main__':
    main()
